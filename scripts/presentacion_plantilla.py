"""Monta cinco diapositivas editables en la plantilla PPTX del estudiante."""
import argparse
import json
from pathlib import Path
import shutil
import sys
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.mathtext import math_to_image
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

REPO=Path(__file__).resolve().parents[1]
BLUE='0760AB';INK='18364B';MUTED='61798B';PALE='EDF4FA';LINE='D9E5EF';WHITE='FFFFFF'


def text(s,x,y,w,h,txt,size=15,color=INK,bold=False,align=PP_ALIGN.LEFT):
    sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    f=sh.text_frame;f.clear();f.word_wrap=True
    f.margin_left=f.margin_right=Inches(.03);f.margin_top=f.margin_bottom=Inches(.02)
    for i,line in enumerate(str(txt).split('\n')):
        p=f.paragraphs[0] if i==0 else f.add_paragraph()
        p.text=line;p.alignment=align;p.space_after=Pt(4)
        for r in p.runs:r.font.name='Arial';r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
    return sh


def box(s,x,y,w,h,fill=PALE,line=None):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill)
    sh.line.fill.background() if line is None else None
    if line:sh.line.color.rgb=RGBColor.from_string(line)
    try:sh.adjustments[0]=.08
    except (IndexError,AttributeError):pass
    return sh


def image_fit(s,p,x,y,w,h):
    from PIL import Image
    with Image.open(p) as im:ratio=im.width/im.height
    if w/h>ratio:nw=h*ratio;nh=h
    else:nw=w;nh=w/ratio
    s.shapes.add_picture(str(p),Inches(x+(w-nw)/2),Inches(y+(h-nh)/2),width=Inches(nw),height=Inches(nh))


def header(s,n,titulo,sub):
    text(s,8.2,.32,4.55,.23,f'BREASTDCEDL  /  ANTONIO AGUILERA  /  {n:02d}–05',9,MUTED,align=PP_ALIGN.RIGHT)
    text(s,.55,1.0,12.1,.55,titulo,29,BLUE,True)
    text(s,.57,1.58,12,.48,sub,13,MUTED)
    text(s,.55,7.16,8,.18,'Aprendizaje Automático · Proyecto individual · Uso educativo, sin validez clínica',8,MUTED)
    text(s,10,7.16,2.75,.18,'30 / 10 / 2026',8,MUTED,align=PP_ALIGN.RIGHT)


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--plantilla',type=Path,required=True)
    p.add_argument('--salida',type=Path,default=REPO/'runs/presentacion_uax')
    p.add_argument('--regenerar',action='store_true',help='Actualizar solo una presentación previamente creada por este script')
    a=p.parse_args();a.salida=a.salida.resolve()
    if a.salida.exists() and (not a.regenerar or not (a.salida/'PRESENTACION_UAX_5_DIAPOSITIVAS.pptx').is_file() or not (a.salida/'LEEME.txt').is_file()):
        p.error('La salida existe; usa otra carpeta o --regenerar para una presentación propia ya creada.')
    a.salida.mkdir(parents=True,exist_ok=True)
    assets=a.salida/'assets';assets.mkdir(exist_ok=True)
    m=json.loads((REPO/'artifacts/dashboard/resumen.json').read_text(encoding='utf-8'));t=m['test']
    ejemplo=m['ejemplos'][0]
    cortes={}
    for nombre in ejemplo['archivos']:
        muestra,fase=nombre.rsplit('_',1)
        cortes.setdefault(muestra,{})[fase.removesuffix('.png')]=nombre
    corte=list(cortes.values())[len(cortes)//2]
    for fase in ['PRE','EARLY','LATE']:
        shutil.copyfile(REPO/'artifacts/dashboard/ejemplos'/ejemplo['id']/corte[fase],assets/f'{fase}.png')
    prs=Presentation(a.plantilla)
    # Conservar portada y layouts originales, eliminar únicamente las tres páginas de ejemplo.
    for sid in list(prs.slides._sldIdLst)[1:]:
        prs.part.drop_rel(sid.rId);prs.slides._sldIdLst.remove(sid)
    s=prs.slides[0]
    for sh in list(s.shapes):
        if sh.shape_type==17:sh._element.getparent().remove(sh._element)
    # 1. Problema clínico dentro de la portada azul original, sin una portada extra.
    text(s,.65,.67,7.0,.32,'CASO 1 · RESONANCIA DE MAMA Y DEEP LEARNING',12,WHITE,True)
    text(s,.65,1.23,6.8,1.2,'Predicción de respuesta\npatológica completa',35,WHITE,True)
    text(s,.65,2.65,6.4,.65,'Una CNN propia para anticipar pCR a partir\nde imágenes PRE, EARLY y LATE pretratamiento.',19,WHITE)
    for i,fase in enumerate(['PRE','EARLY','LATE']):
        im=assets/f'{fase}.png'
        image_fit(s,im,7.55+i*1.45,1.35,1.32,1.32)
        box(s,7.55+i*1.45,2.73,1.32,.34,BLUE)
        text(s,7.55+i*1.45,2.74,1.32,.25,fase,12,WHITE,True,PP_ALIGN.CENTER)
    text(s,.65,3.65,8.1,.88,'pCR = 1: tras el tratamiento no se identifica cáncer invasivo residual en el tejido evaluado, según la definición del estudio.',19,WHITE)
    for i,(etapa,desc) in enumerate([('ANTES','DCE-MRI'),('DESPUÉS','Tratamiento'),('CONFIRMACIÓN','Anatomía patológica')]):
        x=.65+i*2.6
        text(s,x,4.95,2.5,.23,etapa,10,'BBD9F0',True)
        text(s,x,5.23,2.55,.38,desc,17,WHITE)
        if i<2:text(s,x+2.32,5.14,.25,.3,'→',18,WHITE)
    text(s,4.15,6.18,4.7,.65,'Antonio Aguilera · 30/10/2026\nProyecto educativo: no diagnostica ni recomienda tratamientos.',12,WHITE)
    s.notes_slide.notes_text_frame.text='Distinguir respuesta al tratamiento de detectar cáncer. La etiqueta se confirma después por anatomía patológica. No es probabilidad de supervivencia ni recomendación para evitar cirugía.'
    # 2. Datos y validación cruzada.
    s=prs.slides.add_slide(prs.slide_layouts[0]);header(s,2,'La paciente es la unidad estadística','Cortes correlacionados: no equivalen a observaciones independientes. Ningún patient_id se reparte entre train y validación.')
    for x,n,desc,extra in [(.6,'1.097','pacientes de train','10.945 cortes · cinco folds'),(3.48,'176','pacientes de test','1.758 cortes · evaluación histórica')]:
        box(s,x,2.2,2.68,1.15);text(s,x+.15,2.32,2.4,.48,n,31,BLUE,True);text(s,x+.15,2.84,2.4,.25,desc,13);text(s,x+.15,3.1,2.4,.2,extra,9,MUTED)
    text(s,.65,3.65,5.6,.3,'Desbalance en train: 775 no pCR / 322 pCR',16,BLUE,True)
    box(s,.65,4.1,5.42, .4,'D6E5EF');box(s,.65,4.1,5.42*775/1097,.4,BLUE)
    text(s,.8,4.17,3.3,.2,'70,6 % · no pCR',10,WHITE,True);text(s,4.6,4.17,1.3,.2,'29,4 % · pCR',10,BLUE,True)
    text(s,.65,4.73,5.55,.84,'I-SPY1 · I-SPY2 · Duke\nMáscaras en I-SPY, límites tumorales en Duke: distinta selección de cortes y posible sesgo de cohorte.',13,MUTED)
    text(s,7.0,2.15,5.8,.3,'Cinco folds oficiales por paciente',16,BLUE,True)
    for i in range(5):
        y=2.72+i*.47;text(s,7.0,y, .7,.3,f'Fold {i}',10,MUTED)
        for j in range(5):
            x=7.8+j*.92;box(s,x,y,.82,.34,BLUE if i==j else PALE)
            text(s,x+.02,y+.06,.78,.2,'VAL' if i==j else 'TRAIN',9,WHITE if i==j else MUTED,align=PP_ALIGN.CENTER)
    text(s,7.0,5.3,5.8,.55,'La validación privada del profesorado no está incluida.\nCada paciente valida una vez dentro de la CV pública.',12,MUTED)
    box(s,.6,6.15,12.1,.66)
    text(s,.75,6.25,11.8,.43,'PNG gris 8 bits, 256²  →  PRE/EARLY/LATE  →  float32 /255  →  flips conjuntos solo en train\nSin RGB, sin ImageNet y sin normalizar cada fase por separado.',12,BLUE)
    s.notes_slide.notes_text_frame.text='Explicar fuga si se mezclan cortes, por qué accuracy es insuficiente y cómo las tres fases conservan una ventana de intensidades compartida. No confundir z con tiempo de tratamiento.'
    # 3. Arquitectura: todos los bloques y rótulos son formas/texto editables.
    s=prs.slides.add_slide(prs.slide_layouts[0]);header(s,3,'Arquitectura propia: de tres fases a una probabilidad','97.809 parámetros · cuatro bloques Conv3×3 → BatchNorm → ReLU → MaxPool2 · inicialización Kaiming / Xavier')
    etapas=[('ENTRADA','3 × 256 × 256','PRE · EARLY · LATE',None),('BLOQUE 1','16 × 128 × 128','464 parámetros',464),('BLOQUE 2','32 × 64 × 64','4.672 parámetros',4672),('BLOQUE 3','64 × 32 × 32','18.560 parámetros',18560),('BLOQUE 4','128 × 16 × 16','73.984 parámetros',73984),('CABEZA','128 → 1 logit','129 parámetros',129)]
    for i,(nombre,forma,param,_) in enumerate(etapas):
        x=.57+i*2.11;box(s,x,2.35,1.91,1.8,PALE if i<5 else 'D8EAF7',LINE)
        text(s,x+.1,2.48,1.71,.25,nombre,10,BLUE,True,PP_ALIGN.CENTER)
        if 0<i<5:
            for k in range(3):box(s,x+.59+k*.08,2.94-k*.06,.52,.46,['BFD7EA','8AB9D8',BLUE][k])
        elif i==0:image_fit(s,assets/'EARLY.png',x+.55,2.88,.8,.57)
        else:text(s,x+.14,2.96,1.65,.45,'GAP → Dropout',13,BLUE,True,PP_ALIGN.CENTER)
        text(s,x+.08,3.52,1.75,.3,forma,13,INK,True,PP_ALIGN.CENTER)
        text(s,x+.08,3.86,1.75,.2,param,9,MUTED,align=PP_ALIGN.CENTER)
        if i<5:text(s,x+1.91,3.14,.2,.26,'→',14,BLUE)
    formulas=[('bloque',r'$X_b=\mathrm{MaxPool}_2[\mathrm{ReLU}(\mathrm{BN}(W_b*X_{b-1}))]$'),('gap',r'$f_c=\frac{1}{256}\sum_{u=1}^{16}\sum_{v=1}^{16}X_4(c,u,v)$'),('logit',r'$z_i=w^\top\mathrm{Dropout}_{0.3}(f_i)+b,\quad p_i=\sigma(z_i)=\frac{1}{1+e^{-z_i}}$'),('media',r'$\widehat p=\frac{1}{n}\sum_{i=1}^{n}p_i,\quad \widehat y=\mathbf{1}[\widehat p\geq0.33]$')]
    for n,f in formulas:math_to_image(f,str(assets/f'{n}.png'),dpi=300,color='#0760AB')
    image_fit(s,assets/'bloque.png',.75,4.36,6.15,.4)
    text(s,7.0,4.39,5.7,.3,'Conv: padding1, stride1, sin bias; pooling: stride2.',11,MUTED)
    box(s,.6,5.0,3.5,1.65);text(s,.75,5.13,3.2,.2,'1 / GLOBAL AVERAGE POOLING',10,BLUE,True);image_fit(s,assets/'gap.png',.8,5.47,3.05,.55);text(s,.75,6.13,3.2,.32,'128 mapas → vector de 128 valores',10,MUTED)
    box(s,4.32,5.0,4.35,1.65);text(s,4.47,5.13,4.05,.2,'2 / UN LOGIT Y UNA PROBABILIDAD',10,BLUE,True);image_fit(s,assets/'logit.png',4.5,5.5,4.0,.5);text(s,4.47,6.13,4.05,.32,'Dropout desactivado en inferencia: eval()',10,MUTED)
    box(s,8.9,5.0,3.8,1.65);text(s,9.05,5.13,3.5,.2,'3 / MEDIA POR PACIENTE',10,BLUE,True);image_fit(s,assets/'media.png',9.08,5.48,3.4,.56);text(s,9.05,6.13,3.5,.32,'Media de probabilidades, no de logits.',10,MUTED)
    s.notes_slide.notes_text_frame.text='Conv conserva el tamaño por padding1; MaxPool lo divide por2. Parámetros por bloque incluyen 2 variables aprendidas de BN por canal, no sus estadísticas. La capa final tiene128 pesos y1 bias. GAP promedia256 posiciones; BCEWithLogitsLoss se entrena desde logits. En inferencia se aplica sigmoid y se promedian las probabilidades por paciente.'
    # 4. Experimentos y razonamiento de cierre.
    s=prs.slides.add_slide(prs.slide_layouts[0]);header(s,4,'Entrenamiento: comparar antes de elegir','Mismos folds y ambas pérdidas. AUC media ± desviación; diferencias pequeñas no prueban superioridad estadística.')
    rutas=[('Original',REPO/'cv5_v2/analisis'),('Horizonte 60',REPO/'runs/cv5_v3_60ep/analisis'),('Imagen + clínica',REPO/'runs/cv5_clinica_v1/analisis'),('Doble convolución',REPO/'runs/cv5_doble_v1/analisis'),('lr = 0,0003',REPO/'runs/cv5_simple_lr3e4/analisis')]
    table=s.shapes.add_table(6,3, Inches(.6),Inches(2.25),Inches(7.05),Inches(2.28)).table
    table.columns[0].width=Inches(2.8)
    for j,label in enumerate(['Experimento','BCE normal','BCE ponderada']):table.cell(0,j).text=label
    for i,(nombre,carpeta) in enumerate(rutas,1):
        df=pd.read_csv(carpeta/'metricas_por_fold.csv');table.cell(i,0).text=nombre
        for j,perdida in enumerate(['normal','ponderada'],1):
            v=df[df.perdida==perdida].auc;table.cell(i,j).text=f'{v.mean():.3f} ± {v.std():.3f}'.replace('.',',')
    for i,row in enumerate(table.rows):
        row.height=Inches(.38)
        for cell in row.cells:
            cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string(BLUE if i==0 else (PALE if i%2 else WHITE))
            cell.margin_left=cell.margin_right=Inches(.08);cell.margin_top=Inches(.06)
            for p0 in cell.text_frame.paragraphs:
                for rr in p0.runs:rr.font.name='Arial';rr.font.size=Pt(12);rr.font.bold=i==0;rr.font.color.rgb=RGBColor.from_string(WHITE if i==0 else INK)
    h=pd.DataFrame(json.loads((REPO/'runs/cv5_v3_60ep/fold3_semilla42/normal/historial.json').read_text()))
    fig,ax=plt.subplots(figsize=(6,3.4));ax.plot(h.epoca,h.auc_train,label='Train',color='#0760AB',lw=2);ax.plot(h.epoca,h.auc_val,label='Validación',color='#BF6A37',lw=2);ax.axvline(5,color='gray',ls='--');ax.set(xlabel='Época',ylabel='AUC / paciente',ylim=(.4,.8),xlim=(1,17));ax.spines[['top','right']].set_visible(False);ax.legend(frameon=False);fig.tight_layout();fig.savefig(assets/'curva.png',dpi=220);plt.close(fig)
    image_fit(s,assets/'curva.png',7.97,2.2,4.75,2.8)
    text(s,8.07,5.05,4.5,.55,'Fold 3: se conserva época 5.\nEl sobreentrenamiento posterior se muestra.',11,MUTED)
    box(s,.6,4.82,7.05,.96);text(s,.8,4.97,6.65,.66,'Se conserva la CNN sencilla original: las variantes no aportan una mejora consistente. No se elige por el mejor fold ni por volver a mirar test.',15,BLUE,True)
    text(s,.65,6.0,7.0,.68,'Adam · lr 0,001 · batch 16 · wd 0,0001 · dropout 0,3 · semilla 42\nBCEWithLogitsLoss normal / ponderada: α=1 o N₀/N₁ de train.',12,MUTED)
    text(s,8.07,5.95,4.6,.75,'Final: 5 épocas fijas, mediana de [3,5,29,5,6].\nTodo train → media por paciente → umbral 0,33 por F1 macro OOF.',12,BLUE)
    s.notes_slide.notes_text_frame.text='La arquitectura sigue siendo propia; clínica es solo comparación. Explicar faltantes de menopausia y cohortes. Las épocas seleccionadas no son necesariamente mínimas BCE; los umbrales OOF son optimistas. Posteriormente a consultar el test original se exploraron variantes; se conservó el modelo original. Declarar la salvedad de la guía, no afirmar independencia de toda la búsqueda.'
    # 5. Aplicación publicada, resultados y límites, sin una sexta página de gracias.
    s=prs.slides.add_slide(prs.slide_layouts[0]);header(s,5,'Aplicación publicada: demostración y límites','HTML + JavaScript + FastAPI · mismo modelo, normalización, media por paciente y umbral fijo.')
    appshot=REPO/'artifacts/qa_dashboard/dashboard_resultado.png'
    # Crop visual del screenshot, manteniendo visible el visor y el resultado.
    picture=s.shapes.add_picture(str(appshot), Inches(.6), Inches(2.2),width=Inches(6.8),height=Inches(3.62))
    picture.crop_top=.35;picture.crop_bottom=.03
    box(s,7.7,2.2,5.0,1.1)
    text(s,7.9,2.36,4.6,.3,f'Test histórico · AUC {t["auc"]:.3f}'.replace('.',','),22,BLUE,True)
    text(s,7.9,2.83,4.6,.23,f'IC95% [{t["auc_ic95_bootstrap"][0]:.3f}, {t["auc_ic95_bootstrap"][1]:.3f}] · F1macro {t["f1_macro"]:.3f}'.replace('.',','),12,MUTED)
    vals=[t['confusion'][0][0],t['confusion'][0][1],t['confusion'][1][0],t['confusion'][1][1]]
    for i,(v,label) in enumerate(zip(vals,['VN · no pCR correcta','FP · pCR incorrecta','FN · no pCR incorrecta','VP · pCR correcta'])):
        x=7.7+(i%2)*2.6;y=3.55+(i//2)*.79;box(s,x,y,2.4,.67,PALE if i in (0,3) else 'FAF0E7');text(s,x+.15,y+.1,.65,.42,v,23,BLUE,True);text(s,x+.85,y+.14,1.48,.43,label,9,MUTED)
    text(s,7.75,5.28,4.9,.65,f'Sensibilidad {t["sensibilidad"]:.3f} · especificidad {t["especificidad"]:.3f}\nAccuracy {t["accuracy"]:.3f} < baseline {t["baseline_siempre_no_pcr"]["accuracy"]:.3f}'.replace('.',','),12,MUTED)
    url='https://breastdce-antonio.onrender.com'
    sh=text(s,.65,5.98,6.75,.32,url,16,BLUE,True)
    sh.text_frame.paragraphs[0].runs[0].hyperlink.address=url
    text(s,.65,6.37,6.7,.48,'Cinco muestras nuevas sin reentrenar ni reiniciar.\nPNG256², fases validadas, SHA-256, CPU, tiempo y exportación JSON.',11,MUTED)
    text(s,7.75,6.0,4.9,.86,'Límites: sobreajuste, fallo en I-SPY1 y calibración limitada. Test consultado antes de explorar variantes: declarar la cronología. Uso docente, sin validez clínica.',12,INK)
    s.notes_slide.notes_text_frame.text='Abrir la URL pública y mostrar un corte y el registro JSON. No interpretar una muestra como evaluación de toda la paciente. La matriz ordena no pCR/pCR en filas reales y columnas predichas. No ocultar la menor accuracy frente al baseline ni el rendimiento de I-SPY1. Aclarar cronología del test y límites, sin afirmar cumplimiento metodológico pleno.'
    salida=a.salida/'PRESENTACION_UAX_5_DIAPOSITIVAS.pptx';prs.save(salida)
    assert len(Presentation(salida).slides)==5
    (a.salida/'LEEME.txt').write_text('Presentación basada en la plantilla original: portada y layouts/logotipos conservados. Textos, tablas y diagramas son editables; las ecuaciones y capturas son imágenes. Exactamente cinco diapositivas, con URL pública y notas del ponente. El PDF debe exportarse desde PowerPoint para conservar el aspecto. La plantilla original no se modifica.\n',encoding='utf-8')
    print(salida)


if __name__=='__main__':main()
